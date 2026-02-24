"""Customs report parser and import service for bulk knowledge base ingestion.

Parses Vietnamese customs import/export report files (XLS/XLSX) and extracts
verified product-to-HS code mappings into the knowledge base (lookup_records).
"""

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.hs_code import HSCode
from app.models.lookup_record import LookupRecord
from app.repositories.lookup_record_repository import compute_query_hash

logger = logging.getLogger(__name__)

# Column header names used to auto-detect column positions
HS_CODE_HEADER = "Mã HS"
PRODUCT_NAME_HEADER = "Tên hàng"

# Regex to strip dots/spaces from raw HS code strings (e.g., "8507.60.00" -> "85076000")
HS_CODE_CLEAN_RE = re.compile(r"[.\s]")

# Separator used in product names: "CODE#&Description" -> "Description"
PRODUCT_NAME_SEPARATOR = "#&"


@dataclass
class ParsedRow:
    """A single parsed row from a customs report file."""

    product_name: str
    hs_code: str
    row_number: int


@dataclass
class ImportResult:
    """Summary of an import operation."""

    total_rows: int = 0
    records_imported: int = 0
    duplicates_skipped: int = 0
    unmatched_codes: list[dict[str, str | int]] = field(default_factory=list)
    errors: list[dict[str, str | int]] = field(default_factory=list)


class CustomsReportParser:
    """Parses customs report XLS/XLSX files and extracts product-HS code pairs.

    Supports auto-detection of file format (XLS vs XLSX) and column positions
    by scanning for header rows containing "Ma HS" and "Ten hang".
    """

    def __init__(self, file_path: str | Path) -> None:
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"File not found: {self.file_path}")

    def parse(self) -> list[ParsedRow]:
        """Parse the file and return extracted rows.

        Auto-detects format by file extension (.xls vs .xlsx).
        """
        ext = self.file_path.suffix.lower()
        if ext == ".xls":
            return self._parse_xls()
        elif ext == ".xlsx":
            return self._parse_xlsx()
        else:
            raise ValueError(f"Unsupported file format: {ext}. Expected .xls or .xlsx")

    def _parse_xls(self) -> list[ParsedRow]:
        """Parse an XLS file using xlrd."""
        import xlrd

        wb = xlrd.open_workbook(str(self.file_path))
        sheet = wb.sheet_by_index(0)

        # Find header row and column positions
        hs_col, name_col, data_start_row = self._find_columns_xls(sheet)

        rows: list[ParsedRow] = []
        for r in range(data_start_row, sheet.nrows):
            try:
                raw_hs = str(sheet.cell_value(r, hs_col)).strip()
                raw_name = str(sheet.cell_value(r, name_col)).strip()

                if not raw_hs or not raw_name:
                    continue

                hs_code = self._normalize_hs_code(raw_hs)
                if not hs_code:
                    continue

                product_name = self._normalize_product_name(raw_name)
                if not product_name:
                    continue

                rows.append(ParsedRow(
                    product_name=product_name,
                    hs_code=hs_code,
                    row_number=r + 1,  # 1-indexed for user-facing output
                ))
            except Exception as e:
                logger.warning(f"Error parsing row {r + 1}: {e}")
                continue

        return rows

    def _parse_xlsx(self) -> list[ParsedRow]:
        """Parse an XLSX file using openpyxl."""
        from openpyxl import load_workbook

        wb = load_workbook(str(self.file_path), read_only=True, data_only=True)
        try:
            ws = wb.active

            # Find header row and column positions
            hs_col, name_col, data_start_row = self._find_columns_xlsx(ws)

            rows: list[ParsedRow] = []
            # Start iteration from data_start_row + 1 (1-indexed) to skip header rows
            # This avoids re-iterating over rows already scanned during header detection
            for r_idx, row in enumerate(
                ws.iter_rows(min_row=data_start_row + 1, values_only=True)
            ):
                actual_row = data_start_row + 1 + r_idx  # 1-indexed absolute row number

                try:
                    raw_hs = str(row[hs_col]).strip() if row[hs_col] is not None else ""
                    raw_name = str(row[name_col]).strip() if row[name_col] is not None else ""

                    if not raw_hs or not raw_name:
                        continue

                    hs_code = self._normalize_hs_code(raw_hs)
                    if not hs_code:
                        continue

                    product_name = self._normalize_product_name(raw_name)
                    if not product_name:
                        continue

                    rows.append(ParsedRow(
                        product_name=product_name,
                        hs_code=hs_code,
                        row_number=actual_row,
                    ))
                except Exception as e:
                    logger.warning(f"Error parsing row {actual_row}: {e}")
                    continue
        finally:
            wb.close()

        return rows

    def _find_columns_xls(self, sheet: object) -> tuple[int, int, int]:
        """Find HS code and product name column indices in an XLS sheet.

        Scans the first 20 rows for a row containing both "Ma HS" and "Ten hang".

        Args:
            sheet: An xlrd Sheet object.

        Returns:
            Tuple of (hs_code_col, product_name_col, data_start_row)
        """
        for r in range(min(20, sheet.nrows)):
            hs_col = None
            name_col = None
            for c in range(sheet.ncols):
                val = str(sheet.cell_value(r, c)).strip()
                if val == HS_CODE_HEADER:
                    hs_col = c
                elif val == PRODUCT_NAME_HEADER:
                    name_col = c

            if hs_col is not None and name_col is not None:
                logger.info(
                    f"Found headers at row {r}: "
                    f"HS code col={hs_col}, product name col={name_col}"
                )
                return hs_col, name_col, r + 1

        raise ValueError(
            f"Could not find header row with '{HS_CODE_HEADER}' and "
            f"'{PRODUCT_NAME_HEADER}' in {self.file_path}"
        )

    def _find_columns_xlsx(self, ws: object) -> tuple[int, int, int]:
        """Find HS code and product name column indices in an XLSX sheet.

        Returns:
            Tuple of (hs_code_col, product_name_col, data_start_row)
            Column indices are 0-based for use with iter_rows(values_only=True).
        """
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
            actual_row = r_idx + 1  # 1-indexed
            hs_col = None
            name_col = None
            for c_idx, cell_val in enumerate(row):
                val = str(cell_val).strip() if cell_val is not None else ""
                if val == HS_CODE_HEADER:
                    hs_col = c_idx
                elif val == PRODUCT_NAME_HEADER:
                    name_col = c_idx

            if hs_col is not None and name_col is not None:
                logger.info(
                    f"Found headers at row {actual_row}: "
                    f"HS code col={hs_col}, product name col={name_col}"
                )
                return hs_col, name_col, actual_row

        raise ValueError(
            f"Could not find header row with '{HS_CODE_HEADER}' and "
            f"'{PRODUCT_NAME_HEADER}' in {self.file_path}"
        )

    @staticmethod
    def _normalize_hs_code(raw: str) -> str | None:
        """Normalize and validate an HS code string.

        Strips dots, spaces, and trailing '.0' from float conversion.
        Returns the 8-digit code or None if invalid.
        """
        # Handle float-like strings from Excel (e.g., "39269099.0")
        if raw.endswith(".0"):
            raw = raw[:-2]

        # Strip dots and spaces
        cleaned = HS_CODE_CLEAN_RE.sub("", raw)

        # Must be exactly 8 digits
        if len(cleaned) == 8 and cleaned.isdigit():
            return cleaned

        return None

    @staticmethod
    def _normalize_product_name(raw: str) -> str:
        """Normalize a product description.

        - Strips the code prefix before '#&' separator if present
        - Trims whitespace
        - Normalizes Unicode to NFC form
        """
        # Strip code prefix: "CODE#&Description" -> "Description"
        if PRODUCT_NAME_SEPARATOR in raw:
            raw = raw.split(PRODUCT_NAME_SEPARATOR, 1)[1]

        # Trim whitespace and normalize Unicode
        normalized = raw.strip()
        normalized = unicodedata.normalize("NFC", normalized)

        return normalized


class CustomsImportService:
    """Imports parsed customs report rows into the knowledge base.

    Performs bulk HS code lookup, deduplication, and batch insert
    of LookupRecord entries.
    """

    BATCH_SIZE = 500

    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def import_rows(
        self,
        parsed_rows: list[ParsedRow],
        source_file: str,
        company_name: str,
    ) -> ImportResult:
        """Import parsed rows into the knowledge base.

        Args:
            parsed_rows: List of ParsedRow from CustomsReportParser.
            source_file: Original filename for traceability notes.
            company_name: Company name for traceability notes.

        Returns:
            ImportResult with counts and error details.
        """
        result = ImportResult(total_rows=len(parsed_rows))

        if not parsed_rows:
            return result

        # Step 1: Batch lookup all distinct HS codes -> {code: id}
        distinct_codes = {row.hs_code for row in parsed_rows}
        hs_code_map = await self._batch_lookup_hs_codes(distinct_codes)

        # Step 2: Compute query hashes for all rows and batch check for duplicates
        hash_to_row: dict[str, ParsedRow] = {}
        for row in parsed_rows:
            qhash = compute_query_hash(row.product_name)
            # If multiple rows have same hash, keep the first one
            if qhash not in hash_to_row:
                hash_to_row[qhash] = row

        existing_hashes = await self._batch_check_existing_hashes(set(hash_to_row.keys()))

        # Step 3: Build records to insert
        records_to_insert: list[LookupRecord] = []
        seen_hashes: set[str] = set()

        for row in parsed_rows:
            qhash = compute_query_hash(row.product_name)

            # Skip within-file duplicates
            if qhash in seen_hashes:
                result.duplicates_skipped += 1
                continue
            seen_hashes.add(qhash)

            # Skip if already in knowledge base
            if qhash in existing_hashes:
                result.duplicates_skipped += 1
                continue

            # Check if HS code exists in database
            hs_code_id = hs_code_map.get(row.hs_code)
            if hs_code_id is None:
                result.unmatched_codes.append({
                    "code": row.hs_code,
                    "row": row.row_number,
                    "product_name": row.product_name[:100],
                })
                continue

            records_to_insert.append(LookupRecord(
                query_text=row.product_name,
                query_hash=qhash,
                matched_hs_code_id=hs_code_id,
                correct_hs_code_id=hs_code_id,
                is_verified=True,
                search_method="customs_import",
                confidence_score=100,
                notes=f"Source: {source_file} ({company_name})",
            ))

        # Step 4: Batch insert
        for i in range(0, len(records_to_insert), self.BATCH_SIZE):
            batch = records_to_insert[i:i + self.BATCH_SIZE]
            self.session.add_all(batch)
            await self.session.flush()
            logger.info(
                f"Inserted batch {i // self.BATCH_SIZE + 1}: "
                f"{len(batch)} records"
            )

        result.records_imported = len(records_to_insert)

        return result

    async def _batch_lookup_hs_codes(
        self, codes: set[str]
    ) -> dict[str, int]:
        """Look up all HS codes in a single query.

        Returns:
            Mapping of code string -> hs_codes.id
        """
        if not codes:
            return {}

        query = select(HSCode.id, HSCode.code).where(HSCode.code.in_(codes))
        result = await self.session.execute(query)
        return {row.code: row.id for row in result.all()}

    async def _batch_check_existing_hashes(
        self, hashes: set[str]
    ) -> set[str]:
        """Check which query hashes already exist in lookup_records.

        Returns:
            Set of hashes that already exist.
        """
        if not hashes:
            return set()

        # Process in chunks to avoid query parameter limits
        existing: set[str] = set()
        hash_list = list(hashes)
        chunk_size = 1000

        for i in range(0, len(hash_list), chunk_size):
            chunk = hash_list[i:i + chunk_size]
            query = (
                select(LookupRecord.query_hash)
                .where(LookupRecord.query_hash.in_(chunk))
                .distinct()
            )
            result = await self.session.execute(query)
            existing.update(row[0] for row in result.all())

        return existing
